#!/usr/bin/env python3
"""Process every capture note in data/intake/ into a row in data/knowledge_base.xlsx.

Usage:
    python scripts/parse_intake.py

Each intake file is a simple "key: value" list (see data/intake/_template.md). This script:
  1. Parses every *.md file in data/intake/ (skipping _template.md and dotfiles).
  2. Appends one row per file to data/knowledge_base.xlsx, in the column order defined
     in scripts/schema.py.
  3. Moves the processed file into data/processed/ so it isn't parsed twice.

It does not fetch anything from the internet and does not talk to any social media API —
it only structures notes you've already captured by hand.
"""
from __future__ import annotations

import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from schema import AGENCIES, COLUMNS, INTAKES, LEVELS, REGIONS, STATUS_VALUES  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
INTAKE_DIR = ROOT / "data" / "intake"
PROCESSED_DIR = ROOT / "data" / "processed"
WORKBOOK_PATH = ROOT / "data" / "knowledge_base.xlsx"

FIELD_LINE_RE = re.compile(r"^([a-z_]+):\s*(.*)$")


def parse_intake_file(path: Path) -> dict:
    fields: dict[str, str] = {col: "" for col in COLUMNS}
    fields["status"] = "draft"

    text = path.read_text(encoding="utf-8")
    header, _, raw_post = text.partition("--- RAW POST CONTENT ---")

    other_details_extra = []
    for line in header.splitlines():
        line = line.strip()
        if not line or line.startswith("<!--") or line.startswith("-->"):
            continue
        match = FIELD_LINE_RE.match(line)
        if not match:
            continue
        key, value = match.group(1), match.group(2).strip()
        if key in fields:
            fields[key] = value

    if raw_post.strip():
        other_details_extra.append("[raw post excerpt on file: " + path.name + "]")

    if other_details_extra and not fields.get("other_details"):
        fields["other_details"] = " ".join(other_details_extra)

    fields["date_captured"] = fields["date_captured"] or datetime.now().strftime("%Y-%m-%d")
    return fields


def load_or_create_workbook() -> Workbook:
    if WORKBOOK_PATH.exists():
        return load_workbook(WORKBOOK_PATH)

    wb = Workbook()
    ws = wb.active
    ws.title = "offers"
    ws.append(COLUMNS)
    for col_idx, _ in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 22

    _add_dropdown(ws, "source_agency", AGENCIES)
    _add_dropdown(ws, "region", REGIONS)
    _add_dropdown(ws, "intake", INTAKES)
    _add_dropdown(ws, "status", STATUS_VALUES)
    _add_dropdown(ws, "level", LEVELS)
    return wb


def _add_dropdown(ws, column_name: str, values: list[str]) -> None:
    col_idx = COLUMNS.index(column_name) + 1
    col_letter = ws.cell(row=1, column=col_idx).column_letter
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}1048576")


def next_row_id(ws) -> int:
    max_id = 0
    id_col = COLUMNS.index("row_id") + 1
    for row in ws.iter_rows(min_row=2, values_only=False):
        cell = row[id_col - 1]
        if isinstance(cell.value, int):
            max_id = max(max_id, cell.value)
    return max_id + 1


def main() -> None:
    INTAKE_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

    intake_files = sorted(
        p for p in INTAKE_DIR.glob("*.md") if p.name != "_template.md" and not p.name.startswith(".")
    )
    if not intake_files:
        print("No new intake files in data/intake/. Nothing to do.")
        return

    wb = load_or_create_workbook()
    ws = wb["offers"]

    processed = 0
    for path in intake_files:
        fields = parse_intake_file(path)
        fields["row_id"] = next_row_id(ws)
        ws.append([fields.get(col, "") for col in COLUMNS])

        dest = PROCESSED_DIR / path.name
        shutil.move(str(path), str(dest))
        processed += 1
        print(f"  + {path.name} -> row_id {fields['row_id']} ({fields.get('university_name') or 'no university name given'})")

    wb.save(WORKBOOK_PATH)
    print(f"\nProcessed {processed} intake file(s). Workbook saved to {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
